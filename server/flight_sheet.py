# -*- coding: utf-8 -*-
"""Полётная ведомость для учётной записи СИ_Тестер.

Слева пять столбцов из базы курсантов. На каждый день два столбца: Время и Оценка.
Над ними две строки: название дня и дата, оба поля можно править.
Кнопка «+» добавляет следующий день со следующей датой.
"""

import io
import re
from datetime import date, datetime, timedelta, timezone

FLIGHT_SHEET_LOGIN = "си_тестер"
MAX_DAYS = 31

FIXED_COLUMNS = (
    {"id": "npp", "title": "№ п/п", "key": "npp"},
    {"id": "stream", "title": "№ Потока", "key": "stream"},
    {"id": "group", "title": "№ Группы", "key": "group"},
    {"id": "call", "title": "Позывной", "key": "call"},
    {"id": "num", "title": "Номер", "key": "num"},
)

_DAY_KINDS = (
    ("time", "Время"),
    ("grade", "Оценка"),
)

_MONTHS = (
    "янв",
    "фев",
    "мар",
    "апр",
    "май",
    "июн",
    "июл",
    "авг",
    "сен",
    "окт",
    "ноя",
    "дек",
)
_MONTH_INDEX = {name: index + 1 for index, name in enumerate(_MONTHS)}
_MSK = timezone(timedelta(hours=3))

_FIXED_ALIASES = (
    ("npp", ("№ п/п", "n п/п", "п/п", "номер п/п")),
    ("stream", ("№ потока", "номер потока", "поток")),
    ("group", ("№ группы", "номер группы", "группа")),
    ("call", ("позывной курсанта", "позывной")),
    ("num", ("номер курсанта", "номер", "№")),
)


def user_can_use(user) -> bool:
    if not user:
        return False
    for key in ("login", "display_name"):
        if (user.get(key) or "").strip().casefold() == FLIGHT_SHEET_LOGIN:
            return True
    return False


def _norm_header(text: str) -> str:
    s = (text or "").replace("\u00a0", " ").replace("ё", "е")
    s = s.casefold().replace(".", "")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour or value.minute:
            return value.strftime("%d.%m.%Y %H:%M")
        return value.strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).replace("\u00a0", " ").strip()


def _fixed_key(label: str):
    n = _norm_header(label)
    if not n:
        return None
    for key, aliases in _FIXED_ALIASES:
        for alias in aliases:
            if n == _norm_header(alias):
                return key
    return None


def _today() -> date:
    return datetime.now(_MSK).date()


def format_day_date(value: date) -> str:
    return f"{value.day}.{_MONTHS[value.month - 1]}.{value.year % 100:02d}"


def _parse_day_date(text: str):
    raw = _cell_str(text).casefold().replace("ё", "е")
    if not raw:
        return None
    month_name = re.fullmatch(r"(\d{1,2})\.([a-zа-я]+)\.(\d{2,4})", raw)
    if month_name:
        month = _MONTH_INDEX.get(month_name.group(2)) or _MONTH_INDEX.get(month_name.group(2)[:3])
        year = int(month_name.group(3))
        if year < 100:
            year += 2000
        if not month:
            return None
        try:
            return date(year, month, int(month_name.group(1)))
        except ValueError:
            return None
    numeric = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", raw)
    if numeric:
        year = int(numeric.group(3))
        if year < 100:
            year += 2000
        try:
            return date(year, int(numeric.group(2)), int(numeric.group(1)))
        except ValueError:
            return None
    iso = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})", raw)
    if iso:
        try:
            return date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))
        except ValueError:
            return None
    return None


def _next_day_date(previous: str, day_number: int) -> str:
    parsed = _parse_day_date(previous)
    if parsed:
        return format_day_date(parsed + timedelta(days=1))
    return format_day_date(_today() + timedelta(days=day_number - 1))


def _day_kind(label: str):
    n = _norm_header(label)
    n = re.sub(r"\s+\d+$", "", n).strip()
    if n in ("дата", "date"):
        return "date"
    if n in ("время", "time"):
        return "time"
    if n in ("оценка", "grade"):
        return "grade"
    return None


def day_count(columns) -> int:
    found = 0
    for col in columns or []:
        day = int(col.get("day") or 0)
        if day > found:
            found = day
    return found


def is_current_schema(sheet) -> bool:
    columns = (sheet or {}).get("columns") or []
    ids = {col.get("id") for col in columns}
    has_date = any(col.get("key") == "date" for col in columns)
    has_time = any(col.get("key") == "time" for col in columns)
    has_caption = any(col.get("day_date") for col in columns if col.get("day"))
    return "num" in ids and "call" in ids and has_time and has_caption and not has_date and day_count(columns) >= 1


def _day_columns(day: int, title: str, day_date: str):
    cols = []
    for kind, kind_title in _DAY_KINDS:
        cols.append(
            {
                "id": f"d{day}_{kind}",
                "title": kind_title,
                "key": kind,
                "day": day,
                "day_title": title,
                "day_date": day_date,
                "locked": False,
            }
        )
    return cols


def headers_from_columns(columns):
    found = {}
    for col in columns or []:
        day = int(col.get("day") or 0)
        if not day or day in found:
            continue
        found[day] = {
            "day": day,
            "title": col.get("day_title") or f"День {day}",
            "date": col.get("day_date") or "",
        }
    return [found[day] for day in sorted(found)]


def apply_day_headers(columns, headers):
    incoming = {}
    for item in headers or []:
        if not isinstance(item, dict):
            continue
        day = int(item.get("day") or 0)
        if day:
            incoming[day] = item
    updated = []
    for col in columns or []:
        col = dict(col)
        day = int(col.get("day") or 0)
        item = incoming.get(day)
        if item:
            title = _cell_str(item.get("title") if "title" in item else col.get("day_title"))
            day_date = _cell_str(item.get("date") if "date" in item else col.get("day_date"))
            col["day_title"] = title[:80]
            col["day_date"] = day_date[:40]
        updated.append(col)
    return updated


def _columns_for_days(days: int, headers=None):
    by_day = {int(item.get("day") or 0): item for item in (headers or [])}
    start = _today()
    columns = []
    for spec in FIXED_COLUMNS:
        columns.append(
            {
                "id": spec["id"],
                "title": spec["title"],
                "key": spec["key"],
                "day": 0,
                "locked": True,
            }
        )
    for day in range(1, days + 1):
        item = by_day.get(day) or {}
        title = _cell_str(item.get("title") or "") or f"День {day}"
        day_date = _cell_str(item.get("date") or "") or format_day_date(start + timedelta(days=day - 1))
        columns.extend(_day_columns(day, title, day_date))
    return columns


def _cadet_identity(cadet, index: int):
    return {
        "npp": str(index),
        "stream": _cell_str(cadet.get("stream_num") or ""),
        "group": _cell_str(cadet.get("group_num") or ""),
        "call": _cell_str(cadet.get("callsign") or ""),
        "num": _cell_str(cadet.get("student_num") or ""),
    }


def _row_index_by_num(rows):
    found = {}
    for row in rows or []:
        num = _cell_str(row.get("num") or "")
        if num and num not in found:
            found[num] = row
    return found


def _blank_days(columns):
    return {col["id"]: "" for col in columns if not col.get("locked")}


def build_rows(columns, cadets, previous_rows=None):
    previous = _row_index_by_num(previous_rows)
    day_ids = [col["id"] for col in columns if not col.get("locked")]
    rows = []
    for index, cadet in enumerate(cadets or [], start=1):
        identity = _cadet_identity(cadet, index)
        row = {col["id"]: "" for col in columns}
        row.update(identity)
        old = previous.get(identity["num"])
        if old:
            for cid in day_ids:
                row[cid] = _cell_str(old.get(cid) or "")
        rows.append(row)
    return rows


def build_sheet(cadets, days: int = 1, previous_rows=None, source_name: str = "", headers=None):
    days = max(1, min(MAX_DAYS, int(days)))
    columns = _columns_for_days(days, headers)
    rows = build_rows(columns, cadets, previous_rows)
    return {
        "columns": columns,
        "rows": rows,
        "source_name": (source_name or "")[:180],
        "days": days,
        "row_count": len(rows),
    }


def refresh_sheet(sheet, cadets):
    columns = sheet.get("columns") or []
    days = day_count(columns) or 1
    return build_sheet(
        cadets,
        days=days,
        previous_rows=sheet.get("rows") or [],
        source_name=sheet.get("source_name") or "",
        headers=headers_from_columns(columns),
    )


def migrate_sheet(sheet, cadets):
    if is_current_schema(sheet):
        return refresh_sheet(sheet, cadets)
    columns = sheet.get("columns") or []
    days = day_count(columns) or 1
    headers = []
    for day in range(1, days + 1):
        found = ""
        for row in sheet.get("rows") or []:
            parsed = _parse_day_date(row.get(f"d{day}_date") or "")
            if parsed:
                found = format_day_date(parsed)
                break
        headers.append(
            {
                "day": day,
                "title": f"День {day}",
                "date": found or format_day_date(_today() + timedelta(days=day - 1)),
            }
        )
    return build_sheet(
        cadets,
        days=days,
        previous_rows=sheet.get("rows") or [],
        source_name=sheet.get("source_name") or "",
        headers=headers,
    )


def add_day(sheet, cadets):
    columns = apply_day_headers(sheet.get("columns") or [], headers_from_columns(sheet.get("columns") or []))
    days = day_count(columns)
    if days >= MAX_DAYS:
        raise ValueError(f"Больше {MAX_DAYS} дней добавить нельзя")
    headers = headers_from_columns(columns)
    next_day = days + 1 if days else 1
    previous = headers[-1]["date"] if headers else ""
    headers.append(
        {
            "day": next_day,
            "title": f"День {next_day}",
            "date": _next_day_date(previous, next_day),
        }
    )
    return build_sheet(
        cadets,
        days=next_day,
        previous_rows=sheet.get("rows") or [],
        source_name=sheet.get("source_name") or "",
        headers=headers,
    )


def sanitize_rows(columns, rows):
    ids = [c["id"] for c in columns if c.get("id")]
    if not isinstance(rows, list):
        raise ValueError("Нет строк таблицы")
    if len(rows) > 3000:
        raise ValueError("Слишком много строк")
    clean = []
    for raw in rows:
        if not isinstance(raw, dict):
            continue
        row = {}
        for cid in ids:
            value = raw.get(cid, "")
            if value is None:
                value = ""
            row[cid] = _cell_str(value)[:400]
        clean.append(row)
    return clean


def _read_xlsx(raw: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=False)
    try:
        ws = wb.worksheets[0]
        max_row = min(ws.max_row or 0, 3000)
        max_col = min(ws.max_column or 0, 5 + MAX_DAYS * 2)
        grid = []
        for r in range(1, max_row + 1):
            grid.append([_cell_str(ws.cell(r, c).value) for c in range(1, max_col + 1)])
    finally:
        wb.close()
    while grid and not any(grid[-1]):
        grid.pop()
    return grid


def _header_score(row) -> int:
    score = 0
    for cell in row:
        if _fixed_key(cell) or _day_kind(cell):
            score += 1
    return score


def _caption_above(grid, header_at, col_index):
    title = ""
    day_date = ""
    for offset in (1, 2):
        row_at = header_at - offset
        if row_at < 0 or col_index >= len(grid[row_at]):
            continue
        cell = grid[row_at][col_index]
        if not cell or _fixed_key(cell) or _day_kind(cell):
            continue
        if _norm_header(cell) in ("из базы", "курсанты"):
            continue
        parsed = _parse_day_date(cell)
        if parsed and not day_date:
            day_date = format_day_date(parsed)
            continue
        if not title:
            title = cell
    return title, day_date


def _parse_header(grid, header_at):
    row = grid[header_at]
    fixed_at = {}
    day_cells = []
    for index, cell in enumerate(row):
        if not cell:
            continue
        fixed = _fixed_key(cell)
        kind = _day_kind(cell)
        if fixed and fixed not in fixed_at:
            fixed_at[fixed] = index
            continue
        if kind:
            day_cells.append((index, kind))
            continue
        if _norm_header(cell) in ("курсанты", "из базы", "день") or _norm_header(cell).startswith("день "):
            continue
        raise ValueError(
            f"Столбец «{cell}» не из шаблона полётной ведомости. Скачайте шаблон и заполните его."
        )
    if "num" not in fixed_at:
        raise ValueError("В шаблоне нет столбца «Номер»")
    pairs = []
    index = 0
    while index < len(day_cells):
        _col, kind = day_cells[index]
        if kind == "date":
            index += 1
            continue
        if (
            kind == "time"
            and index + 1 < len(day_cells)
            and day_cells[index + 1][1] == "grade"
        ):
            time_at = day_cells[index][0]
            title, day_date = _caption_above(grid, header_at, time_at)
            number = len(pairs) + 1
            pairs.append(
                {
                    "time_at": time_at,
                    "grade_at": day_cells[index + 1][0],
                    "title": title or f"День {number}",
                    "date": day_date or format_day_date(_today() + timedelta(days=number - 1)),
                }
            )
            index += 2
            continue
        raise ValueError("Каждый день — это два столбца подряд: Время и Оценка")
    if not pairs:
        raise ValueError("В шаблоне нет столбцов «Время» и «Оценка»")
    if len(pairs) > MAX_DAYS:
        raise ValueError(f"В шаблоне больше {MAX_DAYS} дней")
    return fixed_at, pairs


def parse_template(raw: bytes, filename: str, cadets):
    name = (filename or "").strip()
    low = name.lower()
    if low.endswith(".doc") or low.endswith(".docx"):
        raise ValueError("Нужен Excel-шаблон полётной ведомости (.xlsx)")
    if not (low.endswith(".xlsx") or low.endswith(".xlsm")):
        raise ValueError("Нужен файл Excel (.xlsx), скачанный как шаблон полётной ведомости")
    grid = _read_xlsx(raw)
    if not grid:
        raise ValueError("В файле нет таблицы")
    header_at = max(range(min(8, len(grid))), key=lambda i: _header_score(grid[i]))
    if _header_score(grid[header_at]) < 2:
        raise ValueError("Не найдена шапка шаблона. Скачайте шаблон полётной ведомости.")
    fixed_at, pairs = _parse_header(grid, header_at)
    num_at = fixed_at["num"]
    uploaded = {}
    skipped = []
    for row in grid[header_at + 1 :]:
        if num_at >= len(row):
            continue
        num = _cell_str(row[num_at])
        if not num:
            continue
        values = []
        for pair in pairs:
            time_value = row[pair["time_at"]] if pair["time_at"] < len(row) else ""
            grade_value = row[pair["grade_at"]] if pair["grade_at"] < len(row) else ""
            values.extend((time_value, grade_value))
        if num in uploaded:
            continue
        uploaded[num] = values
    known = {}
    for cadet in cadets or []:
        num = _cell_str(cadet.get("student_num") or "")
        if num:
            known[num] = cadet
    for num in uploaded:
        if num not in known:
            skipped.append(num)
    days = len(pairs)
    headers = [
        {"day": index, "title": pair["title"], "date": pair["date"]}
        for index, pair in enumerate(pairs, start=1)
    ]
    columns = _columns_for_days(days, headers)
    rows = build_rows(columns, cadets, previous_rows=None)
    day_ids = [col["id"] for col in columns if not col.get("locked")]
    filled = 0
    for row in rows:
        values = uploaded.get(row.get("num") or "")
        if not values:
            continue
        for cid, value in zip(day_ids, values):
            row[cid] = _cell_str(value)[:400]
            if row[cid]:
                filled += 1
    matched = sum(1 for row in rows if uploaded.get(row.get("num") or ""))
    bits = [f"Дней в шаблоне: {days}", f"Строк из базы: {len(rows)}", f"Заполнено ячеек: {filled}"]
    if skipped:
        shown = ", ".join(skipped[:8])
        tail = "" if len(skipped) <= 8 else "…"
        bits.append(f"Номера из файла, которых нет в базе: {shown}{tail}")
    return {
        "columns": columns,
        "rows": rows,
        "source_name": name[:180],
        "days": days,
        "row_count": len(rows),
        "message": ". ".join(bits),
        "matched": matched,
    }


def build_template_bytes(sheet) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    columns = sheet.get("columns") or _columns_for_days(1)
    rows = sheet.get("rows") or []
    headers = headers_from_columns(columns)
    days = len(headers) or 1
    wb = Workbook()
    ws = wb.active
    ws.title = "Полётная ведомость"
    header_font = Font(bold=True, color="001014")
    base_fill = PatternFill("solid", fgColor="19F0F0")
    day_fill = PatternFill("solid", fgColor="8FD6D8")
    input_fill = PatternFill("solid", fgColor="FFF8E6")
    lock_fill = PatternFill("solid", fgColor="E8F7F7")
    thin = Border(
        left=Side(style="thin", color="0B7F86"),
        right=Side(style="thin", color="0B7F86"),
        top=Side(style="thin", color="0B7F86"),
        bottom=Side(style="thin", color="0B7F86"),
    )
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=5)
    label = ws.cell(1, 1, "Из базы")
    label.font = header_font
    label.fill = base_fill
    label.alignment = center
    for col_index in range(1, 6):
        for row_index in (1, 2):
            cell = ws.cell(row_index, col_index)
            cell.fill = base_fill
            cell.border = thin
            cell.alignment = center
    for item in headers:
        day = int(item["day"])
        start = 6 + (day - 1) * 2
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=start + 1)
        ws.merge_cells(start_row=2, start_column=start, end_row=2, end_column=start + 1)
        title_cell = ws.cell(1, start, item.get("title") or f"День {day}")
        date_cell = ws.cell(2, start, item.get("date") or "")
        for cell in (title_cell, date_cell):
            cell.font = header_font
            cell.fill = day_fill
            cell.alignment = center
        for offset in (0, 1):
            for row_index in (1, 2):
                extra = ws.cell(row_index, start + offset)
                extra.fill = day_fill
                extra.border = thin
                extra.alignment = center
    for col_index, col in enumerate(columns, start=1):
        cell = ws.cell(3, col_index, col["title"])
        cell.font = header_font
        cell.fill = base_fill if col.get("locked") else day_fill
        cell.alignment = center
        cell.border = thin
    for row_index, row in enumerate(rows, start=4):
        for col_index, col in enumerate(columns, start=1):
            cell = ws.cell(row_index, col_index, row.get(col["id"]) or "")
            cell.border = thin
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = lock_fill if col.get("locked") else input_fill
    widths = {"npp": 10, "stream": 14, "group": 14, "call": 18, "num": 12, "time": 14, "grade": 14}
    for col_index, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(col_index)].width = widths.get(col.get("key"), 14)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    ws.freeze_panes = "F4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{max(3, len(rows) + 3)}"

    note = wb.create_sheet("Инструкция")
    note["A1"] = (
        "Столбцы «№ п/п», «№ Потока», «№ Группы», «Позывной» и «Номер» заполнены из базы курсантов. Их не меняйте."
    )
    note["A2"] = "Над каждым днём две строки: название дня и дата. Ниже — «Время» и «Оценка»."
    note["A3"] = "Чтобы добавить день, на платформе нажмите «+» и скачайте шаблон заново."
    note.column_dimensions["A"].width = 110

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
