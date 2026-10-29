# -*- coding: utf-8 -*-
import io
import json
import os
import re
from typing import Optional, List, Any

import qrcode
from qrcode.image.svg import SvgPathImage
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Header, Depends, Query, Cookie, Request
from fastapi.responses import FileResponse, HTMLResponse, Response, JSONResponse, RedirectResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

from . import db
from . import auth
from . import flight_sheet
from . import features
from . import library
from . import static_version
from .inject import inject_bridge

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
STATIC_DIR = os.path.join(os.path.dirname(__file__), "static")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "admin")
MASTER_PASSWORD = os.environ.get("MASTER_PASSWORD", ADMIN_PASSWORD)

app = FastAPI(title="ЛИС · Сервер тестов", version="1.2")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def _startup():
    db.init_db()
    auth.ensure_admin_password(ADMIN_PASSWORD)
    try:
        admin = auth.get_user_by_login("admin")
        if admin:
            with db.connect() as conn:
                conn.execute(
                    "UPDATE tests SET owner_user_id=? WHERE owner_user_id IS NULL OR owner_user_id=''",
                    (admin["id"],),
                )
                conn.execute(
                    "UPDATE cadets SET owner_user_id=? WHERE owner_user_id IS NULL OR owner_user_id=''",
                    (admin["id"],),
                )
    except Exception:
        pass
    try:
        for t in db.list_tests():
            path = os.path.join(db.TESTS_DIR, f"{t['id']}.html")
            if not os.path.isfile(path):
                continue
            with open(path, "r", encoding="utf-8") as f:
                html = f.read()
            html2 = inject_bridge(html, t["id"])
            if html2 != html:
                with open(path, "w", encoding="utf-8") as f:
                    f.write(html2)
            if not t.get("questions_count"):
                db.refresh_questions_count(t["id"], html2)
    except Exception:
        pass
    try:
        db.backfill_questions_counts()
    except Exception:
        pass


SESSION_COOKIE = "lis_session"


def _extract_token(
    authorization: Optional[str] = Header(None),
    x_admin_token: Optional[str] = Header(None),
    lis_session: Optional[str] = Cookie(None, alias=SESSION_COOKIE),
) -> Optional[str]:
    token = None
    if authorization and authorization.lower().startswith("bearer "):
        token = authorization[7:].strip()
    if x_admin_token:
        token = x_admin_token
    if not token and lis_session:
        token = lis_session.strip() or None
    return token or None


def _set_session_cookie(resp: Response, token: str) -> None:
    resp.set_cookie(
        key=SESSION_COOKIE,
        value=token,
        httponly=True,
        samesite="lax",
        path="/",
        max_age=60 * 60 * 24 * 30,
    )


def _clear_session_cookie(resp: Response) -> None:
    resp.delete_cookie(key=SESSION_COOKIE, path="/")


def require_user(
    authorization: Optional[str] = Header(None),
    x_admin_token: Optional[str] = Header(None),
    lis_session: Optional[str] = Cookie(None, alias=SESSION_COOKIE),
) -> dict:
    token = _extract_token(authorization, x_admin_token, lis_session)
    user = auth.user_from_token(token) if token else None
    if not user:
        raise HTTPException(401, "Требуется вход")
    return user


def _user_from_request(request: Request) -> Optional[dict]:
    token = (request.cookies.get(SESSION_COOKIE) or "").strip()
    return auth.user_from_token(token) if token else None


def _guard_html(request: Request) -> Optional[RedirectResponse]:
    """Редирект на /login, если нет валидной сессии."""
    if _user_from_request(request):
        return None
    return RedirectResponse(url="/login", status_code=303)


@app.middleware("http")
async def _block_raw_static_html(request: Request, call_next):
    """Не отдавать HTML из /static без сессии (обход через прямые URL)."""
    path = request.url.path or ""
    if path.startswith("/static/") and path.lower().endswith(".html"):
        low = path.lower()
        if low.endswith("/login.html"):
            return RedirectResponse(url="/login", status_code=303)
        if not _user_from_request(request):
            return RedirectResponse(url="/login", status_code=303)
        if low.endswith("/hub.html"):
            return RedirectResponse(url="/", status_code=303)
        if low.endswith("/admin.html"):
            return RedirectResponse(url="/admin", status_code=303)
        if low.endswith("/workload.html"):
            return RedirectResponse(url="/tools/workload", status_code=303)
    return await call_next(request)


def require_admin_role(user: dict = Depends(require_user)) -> dict:
    if user.get("role") != "admin":
        raise HTTPException(403, "Только для администратора")
    return user


def require_test_access(test_id: str, user: dict) -> dict:
    t = db.get_test(test_id)
    if not t:
        raise HTTPException(404, "Тест не найден")
    if not auth.can_access_test(user, t):
        raise HTTPException(403, "Нет доступа к этому тесту")
    return t


class LoginIn(BaseModel):
    login: str = ""
    password: str


class SubmitIn(BaseModel):
    num: str = Field(..., min_length=1)
    call: str = Field(..., min_length=1)
    score: int = 0
    total: int = 0
    percent: int = 0
    grade: str = ""
    answers: List[Any] = []
    forced: bool = False
    duration_sec: Optional[int] = None
    is_continue: bool = False
    submit_id: str = ""
    start_id: str = ""


class BeginIn(BaseModel):
    num: str = Field(..., min_length=1)
    call: str = ""
    start_id: str = ""


class ContinueIn(BaseModel):
    minutes: int = Field(..., ge=1, le=600)


class GradeIn(BaseModel):
    grade: str = Field(..., min_length=1)


class GradingRuleIn(BaseModel):
    min: int = Field(..., ge=0, le=100)
    grade: str = Field(..., min_length=1)


class GradingIn(BaseModel):
    rules: List[GradingRuleIn]


class UnlockIn(BaseModel):
    test_id: str
    num: str
    call: str = ""
    issue_invite: bool = True


class RedeemIn(BaseModel):
    code: str = ""
    master_password: str = ""
    num: str = ""


class RosterEntryIn(BaseModel):
    num: str = Field(..., min_length=1)
    call: str = ""
    fio: str = ""
    full_name: str = ""
    group_num: str = ""
    stream_num: str = ""
    group: str = ""
    stream: str = ""


class RosterIn(BaseModel):
    text: str = ""
    numbers: List[str] = []
    entries: List[RosterEntryIn] = []


class RetakeIn(BaseModel):
    num: str


class TimerIn(BaseModel):
    minutes: int = Field(..., ge=1, le=600)


class CadetIn(BaseModel):
    callsign: str = ""
    student_num: str = ""
    full_name: str = ""
    group_num: str = ""
    stream_num: str = ""


class CadetBulkDeleteIn(BaseModel):
    ids: List[str] = []


class RosterBulkDeleteIn(BaseModel):
    nums: List[str] = []


# ---------- Auth ----------
@app.post("/api/admin/login")
def admin_login(body: LoginIn):
    try:
        result = auth.login(body.login or "admin", body.password)
    except ValueError as e:
        raise HTTPException(403, str(e))
    resp = JSONResponse(result)
    _set_session_cookie(resp, result["token"])
    return resp


@app.get("/api/admin/me")
def admin_me(user: dict = Depends(require_user)):
    return {"ok": True, "user": user}


@app.get("/api/grading")
def api_get_grading():
    """Общая шкала. Читается тестом курсанта без входа в кабинет."""
    return {"ok": True, "rules": db.get_grading()}


@app.put("/api/admin/grading")
def api_set_grading(body: GradingIn, user: dict = Depends(require_user)):
    if user.get("role") not in ("admin", "senior"):
        raise HTTPException(403, "Шкалу меняют администратор и старший инструктор")
    try:
        rules = db.set_grading([{"min": r.min, "grade": r.grade} for r in body.rules])
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "rules": rules}


@app.post("/api/admin/logout")
def admin_logout(
    authorization: Optional[str] = Header(None),
    x_admin_token: Optional[str] = Header(None),
    lis_session: Optional[str] = Cookie(None, alias=SESSION_COOKIE),
    user: dict = Depends(require_user),
):
    token = _extract_token(authorization, x_admin_token, lis_session)
    if token:
        auth.revoke_token(token)
    resp = JSONResponse({"ok": True})
    _clear_session_cookie(resp)
    return resp


class ChangePasswordIn(BaseModel):
    old_password: str = ""
    new_password: str


@app.post("/api/admin/change-password")
def admin_change_password(body: ChangePasswordIn, user: dict = Depends(require_user)):
    try:
        auth.change_password(user["id"], body.old_password, body.new_password)
    except ValueError as e:
        raise HTTPException(400, str(e))
    # новый токен после смены пароля
    token = auth.create_token(user["id"])
    resp = JSONResponse(
        {"ok": True, "token": token, "user": auth.get_user(user["id"])}
    )
    _set_session_cookie(resp, token)
    return resp


# ---------- Departments / staff / users (admin) ----------
class DepartmentIn(BaseModel):
    name: str


class StaffIn(BaseModel):
    department_id: str
    role: str
    fio: str
    callsign: str = ""


class UserCreateIn(BaseModel):
    login: str
    password: str = ""
    role: str
    department_id: Optional[str] = None
    staff_id: Optional[str] = None
    display_name: str = ""
    must_change_password: bool = True


class UserPasswordIn(BaseModel):
    password: str = ""
    must_change_password: bool = True


class UserUpdateIn(BaseModel):
    login: str
    display_name: str = ""
    role: str
    department_id: Optional[str] = None
    callsign: str = ""
    password: str = ""
    must_change_password: bool = True
    is_tester: bool = False
    # совместимость со старым полем UI
    can_flight_sheet: Optional[bool] = None


@app.get("/api/admin/departments")
def api_list_departments(user: dict = Depends(require_user)):
    # admin / senior — все отделы (база, которую вводит администратор);
    # остальные — только свой отдел
    deps = auth.list_departments()
    if user.get("role") in ("admin", "senior"):
        return deps
    return [d for d in deps if d["id"] == user.get("department_id")]


@app.post("/api/admin/departments")
def api_create_department(body: DepartmentIn, _: dict = Depends(require_admin_role)):
    try:
        return auth.create_department(body.name)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.put("/api/admin/departments/{department_id}")
def api_update_department(
    department_id: str, body: DepartmentIn, _: dict = Depends(require_admin_role)
):
    try:
        return auth.update_department(department_id, body.name)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.delete("/api/admin/departments/{department_id}")
def api_delete_department(department_id: str, _: dict = Depends(require_admin_role)):
    try:
        auth.delete_department(department_id)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True}


@app.get("/api/admin/staff")
def api_list_staff(
    department_id: Optional[str] = None,
    user: dict = Depends(require_user),
):
    # admin / senior — весь ЛС (или фильтр по department_id);
    # остальные — только свой отдел
    if user.get("role") not in ("admin", "senior"):
        department_id = user.get("department_id")
    return auth.list_staff(department_id)


@app.post("/api/admin/staff")
def api_add_staff(body: StaffIn, _: dict = Depends(require_admin_role)):
    try:
        return auth.add_staff(body.department_id, body.role, body.fio, body.callsign)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.put("/api/admin/staff/{staff_id}")
def api_update_staff(
    staff_id: str, body: StaffIn, _: dict = Depends(require_admin_role)
):
    try:
        return auth.update_staff(
            staff_id,
            department_id=body.department_id,
            role=body.role,
            fio=body.fio,
            callsign=body.callsign,
        )
    except ValueError as e:
        raise HTTPException(400, str(e))


def _norm_staff_role(raw: str) -> Optional[str]:
    s = (raw or "").strip().casefold().replace("ё", "е")
    if not s:
        return None
    aliases = {
        "senior": "senior",
        "си": "senior",
        "старший": "senior",
        "старший инструктор": "senior",
        "instructor": "instructor",
        "и": "instructor",
        "инструктор": "instructor",
        "assistant": "assistant",
        "аи": "assistant",
        "ассистент": "assistant",
        "ассистент инструктора": "assistant",
    }
    return aliases.get(s)


def _build_staff_template_bytes():
    """Шаблон личного состава. Колонки только здесь, к шаблону курсантов не относится."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Сотрудники"
    headers = ("Отдел", "Должность", "ФИО", "Позывной")
    ws.append(headers)
    header_font = Font(bold=True, color="001014")
    fill = PatternFill("solid", fgColor="19F0F0")
    thin = Border(
        left=Side(style="thin", color="0B7F86"),
        right=Side(style="thin", color="0B7F86"),
        top=Side(style="thin", color="0B7F86"),
        bottom=Side(style="thin", color="0B7F86"),
    )
    for col, _ in enumerate(headers, 1):
        cell = ws.cell(1, col)
        cell.font = header_font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
    for dep, role, fio, call in (
        ("17", "Старший инструктор", "Иванов Иван Иванович", "Сокол"),
        ("17", "Инструктор", "Петров Пётр Петрович", "Ястреб"),
        ("17", "Ассистент инструктора", "Сидоров Сидор Сидорович", "Беркут"),
        ("1", "СИ", "Пример Старший", ""),
        ("2", "И", "Пример Инструктор", ""),
        ("10", "АИ", "Пример Ассистент", ""),
    ):
        ws.append((dep, role, fio, call))
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 36
    ws.column_dimensions["D"].width = 16
    ws2 = wb.create_sheet("Инструкция")
    ws2["A1"] = "Лист «Сотрудники»: A — номер/название отдела, B — должность, C — ФИО, D — позывной."
    ws2["A2"] = (
        "Отдел: укажите номер или название ровно как в кабинете (например 17). "
        "При загрузке отдел берётся из этой колонки для каждой строки."
    )
    ws2["A3"] = (
        "Должность: Старший инструктор (СИ), Инструктор (И), "
        "Ассистент инструктора (АИ)."
    )
    ws2["A4"] = "Строки-примеры замените своими данными. Сначала создайте отделы в кабинете."
    ws2["A5"] = "Если колонки «Отдел» нет (старый шаблон) — используется отдел из фильтра в «Детализация по отделам»."
    ws2.column_dimensions["A"].width = 110
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _parse_staff_excel(raw: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    ws = wb.active
    entries = []
    # "dep" = Отдел|Должность|ФИО|Позывной; "old" = Должность|ФИО|Позывной
    layout = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if not row:
            continue
        cells = [_cell_str(c) for c in row[:4]]
        while len(cells) < 4:
            cells.append("")
        a, b, c, d = cells
        if not a and not b and not c and not d:
            continue
        s0, s1 = a.casefold(), b.casefold()
        if i == 1 and (
            s0 in ("отдел", "номер отдела", "№ отдела", "department", "dep")
            or s0 in ("должность", "роль", "role", "си/и/аи")
            or s1 in ("должность", "фио", "ф.и.о.", "fio")
        ):
            layout = (
                "dep"
                if s0 in ("отдел", "номер отдела", "№ отдела", "department", "dep")
                else "old"
            )
            continue

        if layout is None:
            if _norm_staff_role(b) and a:
                layout = "dep"
            else:
                layout = "old"

        if layout == "dep":
            dep_name, role_raw, fio, call = a, b, c, d
        else:
            dep_name, role_raw, fio, call = "", a, b, c

        role = _norm_staff_role(role_raw)
        if not role or not fio:
            continue
        entries.append(
            {
                "department_name": dep_name,
                "role": role,
                "fio": fio,
                "callsign": call,
            }
        )
    wb.close()
    return entries


def _resolve_department_id(name: str, deps: list) -> Optional[str]:
    key = (name or "").strip().casefold()
    if not key:
        return None
    for d in deps:
        if (d.get("name") or "").strip().casefold() == key:
            return d["id"]
    # «17.0» из Excel → «17»
    if key.endswith(".0") and key[:-2].isdigit():
        key2 = key[:-2]
        for d in deps:
            if (d.get("name") or "").strip().casefold() == key2:
                return d["id"]
    return None


@app.get("/api/admin/staff-template.xlsx")
def api_staff_template(_: dict = Depends(require_admin_role)):
    data = _build_staff_template_bytes()
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="lis-staff-template.xlsx"'
        },
    )


def _staff_fio_key(fio: str) -> str:
    return db._norm_fio(fio).casefold()


@app.post("/api/admin/staff/import")
async def api_import_staff(
    file: UploadFile = File(...),
    department_id: str = Form(""),
    _: dict = Depends(require_admin_role),
):
    fallback_dep = (department_id or "").strip()
    name = (file.filename or "").lower()
    if not (name.endswith(".xlsx") or name.endswith(".xlsm")):
        raise HTTPException(400, "Нужен файл Excel (.xlsx)")
    raw = await file.read()
    try:
        entries = _parse_staff_excel(raw)
    except Exception as e:
        raise HTTPException(400, "Не удалось прочитать Excel: " + str(e))
    if not entries:
        raise HTTPException(
            400,
            "В файле нет строк с должностью и ФИО. Скачайте шаблон и заполните лист «Сотрудники».",
        )
    deps = auth.list_departments()
    existing = auth.list_staff()
    existing_fio = {_staff_fio_key(s.get("fio") or "") for s in existing if (s.get("fio") or "").strip()}
    seen_fio = set()
    created = []
    errors = []
    duplicates = []
    for idx, e in enumerate(entries, start=1):
        fio = db._norm_fio(e.get("fio") or "")
        fio_key = _staff_fio_key(fio)
        role_label = auth.ROLE_LABELS.get(e.get("role") or "", e.get("role") or "")
        call = (e.get("callsign") or "").strip()
        dep_label = (e.get("department_name") or "").strip() or "—"
        if fio_key and fio_key in existing_fio:
            duplicates.append(
                {
                    "row": idx,
                    "fio": fio,
                    "role": role_label,
                    "callsign": call,
                    "department": dep_label,
                    "reason": "уже есть в таблице",
                }
            )
            continue
        if fio_key and fio_key in seen_fio:
            duplicates.append(
                {
                    "row": idx,
                    "fio": fio,
                    "role": role_label,
                    "callsign": call,
                    "department": dep_label,
                    "reason": "повтор в файле",
                }
            )
            continue
        dep_id = _resolve_department_id(e.get("department_name") or "", deps) or fallback_dep
        if not dep_id:
            errors.append(
                f"Строка {idx}: отдел «{dep_label}» не найден. Создайте отдел или укажите его в фильтре."
            )
            continue
        try:
            created.append(
                auth.add_staff(dep_id, e["role"], fio, call)
            )
            if fio_key:
                seen_fio.add(fio_key)
                existing_fio.add(fio_key)
        except ValueError as err:
            errors.append(f"Строка {idx}: {err}")
    if not created and errors and not duplicates:
        raise HTTPException(400, "; ".join(errors[:5]))
    return {
        "ok": True,
        "imported": len(created),
        "skipped": len(errors) + len(duplicates),
        "errors": errors[:10],
        "duplicates": duplicates,
        "staff": auth.list_staff(fallback_dep or None),
    }


@app.delete("/api/admin/staff/{staff_id}")
def api_delete_staff(staff_id: str, _: dict = Depends(require_admin_role)):
    auth.delete_staff(staff_id)
    return {"ok": True}


@app.get("/api/admin/users")
def api_list_users(_: dict = Depends(require_admin_role)):
    # пароли только здесь и только для роли admin (require_admin_role)
    return auth.list_users(include_passwords=True)


@app.post("/api/admin/users")
def api_create_user(body: UserCreateIn, actor: dict = Depends(require_admin_role)):
    pwd = (body.password or "").strip() or auth.gen_temp_password()
    try:
        user = auth.create_user(
            login_name=body.login,
            password=pwd,
            role=body.role,
            department_id=body.department_id,
            staff_id=body.staff_id,
            display_name=body.display_name,
            must_change_password=body.must_change_password,
            created_by=actor["id"],
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "user": user, "temp_password": pwd}


@app.post("/api/admin/users/from-staff")
def api_create_users_from_staff(actor: dict = Depends(require_admin_role)):
    """Массовое создание учёток по таблице личного состава."""
    return auth.create_users_from_all_staff(created_by=actor["id"])


@app.put("/api/admin/users/{user_id}")
def api_update_user(
    user_id: str, body: UserUpdateIn, actor: dict = Depends(require_admin_role)
):
    try:
        user = auth.update_user(
            user_id,
            login_name=body.login,
            display_name=body.display_name,
            role=body.role,
            department_id=body.department_id,
            callsign=body.callsign,
            password=body.password if (body.password or "").strip() else None,
            must_change_password=body.must_change_password,
            is_tester=body.is_tester
            if body.can_flight_sheet is None
            else body.can_flight_sheet,
            actor_id=actor["id"],
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "user": user}


@app.post("/api/admin/users/{user_id}/password")
def api_reset_user_password(
    user_id: str, body: UserPasswordIn, _: dict = Depends(require_admin_role)
):
    pwd = (body.password or "").strip() or auth.gen_temp_password()
    try:
        auth.set_password(user_id, pwd, must_change=body.must_change_password)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "temp_password": pwd}


class UserActiveIn(BaseModel):
    active: bool = True


@app.post("/api/admin/users/{user_id}/active")
def api_set_user_active(
    user_id: str, body: UserActiveIn, _: dict = Depends(require_admin_role)
):
    auth.set_user_active(user_id, body.active)
    return {"ok": True}


@app.delete("/api/admin/users/{user_id}")
def api_delete_user(user_id: str, actor: dict = Depends(require_admin_role)):
    try:
        auth.delete_user(user_id, actor["id"])
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True}


# ---------- QR ----------
@app.get("/api/qr.png")
@app.get("/api/qr.svg")
def qr_image(data: str = Query(..., min_length=1, max_length=2000)):
    img = qrcode.make(data, image_factory=SvgPathImage, border=2)
    buf = io.BytesIO()
    img.save(buf)
    return Response(buf.getvalue(), media_type="image/svg+xml")


# ---------- Tests ----------
@app.get("/api/tests")
def api_list_tests(user: dict = Depends(require_user)):
    include_pw = user.get("role") == "admin"
    return auth.filter_tests_for_user(
        user, db.list_tests(include_master_password=include_pw)
    )


def _can_publish_for_all(user: dict) -> bool:
    """Галочка/кнопка публикации — не для инструктора и ассистента."""
    return user.get("role") in ("admin", "senior")


def _parse_bool_form(val: Optional[str]) -> bool:
    if val is None:
        return False
    return str(val).strip().lower() in ("1", "true", "yes", "on")


@app.post("/api/tests/upload")
async def api_upload_test(
    file: UploadFile = File(...),
    title: Optional[str] = Form(None),
    timer_minutes: Optional[str] = Form(None),
    master_password: Optional[str] = Form(None),
    published_for_all: Optional[str] = Form(None),
    user: dict = Depends(require_user),
):
    raw = await file.read()
    try:
        html = raw.decode("utf-8")
    except UnicodeDecodeError:
        html = raw.decode("cp1251", errors="replace")
    if "<html" not in html.lower():
        raise HTTPException(400, "Нужен HTML-файл теста")

    name = title or (file.filename or "Тест").rsplit(".", 1)[0]
    m = re.search(r"<title>(.*?)</title>", html, re.I | re.S)
    if m and not title:
        name = re.sub(r"\s*[—\-].*$", "", m.group(1)).strip() or name

    mins = 40
    if timer_minutes is not None and str(timer_minutes).strip() != "":
        try:
            mins = int(str(timer_minutes).strip())
        except ValueError:
            raise HTTPException(400, "Таймер: укажите целое число минут")
        if mins < 1 or mins > 600:
            raise HTTPException(400, "Таймер: от 1 до 600 минут")

    master = (master_password or "").strip()
    want_pub = _parse_bool_form(published_for_all)
    if want_pub and not _can_publish_for_all(user):
        raise HTTPException(
            403,
            "Публиковать тест для всех могут только администратор и старший инструктор",
        )
    tid = db.create_test(
        name,
        require_roster=True,
        timer_minutes=mins,
        owner_user_id=user["id"],
        department_id=user.get("department_id"),
        master_password=master,
        published_for_all=want_pub,
    )
    html = inject_bridge(html, tid)
    path = os.path.join(db.TESTS_DIR, f"{tid}.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    qn = db.refresh_questions_count(tid, html)
    return {
        "id": tid,
        "title": name,
        "link": f"/t/{tid}",
        "require_roster": True,
        "timer_minutes": mins,
        "questions_count": qn,
        "has_master_password": bool(master),
        "published_for_all": want_pub,
        "absolute_hint": f"http://<IP>:8080/t/{tid}",
    }


@app.get("/api/tests/catalog")
def api_tests_catalog(user: dict = Depends(require_user)):
    """Каталог тестов старших инструкторов, опубликованных для всех."""
    return db.list_published_catalog()


@app.post("/api/tests/{test_id}/adopt")
def api_adopt_test(test_id: str, user: dict = Depends(require_user)):
    """Скопировать опубликованный тест себе (своя ссылка, свой список курсантов)."""
    src = db.get_test(test_id)
    if not src:
        raise HTTPException(404, "Тест не найден")
    if not src.get("published_for_all"):
        raise HTTPException(403, "Тест не опубликован для общего пользования")
    owner = auth.get_user(src.get("owner_user_id")) if src.get("owner_user_id") else None
    if not owner or owner.get("role") != "senior":
        raise HTTPException(403, "В каталоге только тесты старших инструкторов")

    src_path = os.path.join(db.TESTS_DIR, f"{test_id}.html")
    if not os.path.isfile(src_path):
        raise HTTPException(404, "Файл теста не найден")
    with open(src_path, "r", encoding="utf-8") as f:
        html = f.read()

    title = (src.get("title") or "Тест").strip()
    mins = src.get("timer_minutes") or 40
    qn = src.get("questions_count") or db.questions_count_from_html(html)
    tid = db.create_test(
        title,
        require_roster=True,
        timer_minutes=mins,
        owner_user_id=user["id"],
        department_id=user.get("department_id"),
        master_password="",
        published_for_all=False,
        questions_count=qn,
    )
    html = inject_bridge(html, tid)
    path = os.path.join(db.TESTS_DIR, f"{tid}.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    qn = db.refresh_questions_count(tid, html) or qn
    return {
        "id": tid,
        "title": title,
        "link": f"/t/{tid}",
        "require_roster": True,
        "timer_minutes": mins,
        "questions_count": qn,
        "published_for_all": False,
        "source_id": test_id,
    }


class TestUpdateIn(BaseModel):
    title: str = Field(..., min_length=1)
    timer_minutes: Optional[int] = Field(None, ge=1, le=600)
    master_password: Optional[str] = None
    published_for_all: Optional[bool] = None


def _is_test_owner(t: dict, user: dict) -> bool:
    return bool(t and user and t.get("owner_user_id") and t.get("owner_user_id") == user.get("id"))


def _can_edit_test_full(t: dict, user: dict) -> bool:
    """Полное редактирование (параметры + банк): владелец или администратор."""
    if not t or not user:
        return False
    if user.get("role") == "admin":
        return True
    return _is_test_owner(t, user)


def _require_test_editor(test_id: str, user: dict) -> dict:
    """Владелец теста или админ; старший — параметры тестов своего отдела."""
    t = require_test_access(test_id, user)
    if _can_edit_test_full(t, user):
        return t
    if user.get("role") == "senior":
        return t
    raise HTTPException(403, "Редактирование этого теста недоступно")


def _require_test_full_editor(test_id: str, user: dict) -> dict:
    """Замена HTML/банка — только владелец или администратор."""
    t = require_test_access(test_id, user)
    if not _can_edit_test_full(t, user):
        raise HTTPException(403, "Менять банк вопросов может только автор теста")
    return t


@app.get("/api/admin/tests/{test_id}")
def api_get_test_admin(test_id: str, user: dict = Depends(require_user)):
    _require_test_editor(test_id, user)
    include_pw = user.get("role") == "admin"
    t = db.get_test(test_id, include_master_password=include_pw)
    if not t:
        raise HTTPException(404, "Тест не найден")
    return t


@app.put("/api/admin/tests/{test_id}")
def api_update_test_admin(
    test_id: str, body: TestUpdateIn, user: dict = Depends(require_user)
):
    _require_test_editor(test_id, user)
    if body.published_for_all is not None and not _can_publish_for_all(user):
        raise HTTPException(403, "Публиковать тест для всех нельзя этой ролью")
    kwargs = {
        "title": body.title,
        "timer_minutes": body.timer_minutes,
    }
    if user.get("role") == "admin" and body.master_password is not None:
        kwargs["master_password"] = body.master_password
    if body.published_for_all is not None and _can_publish_for_all(user):
        kwargs["published_for_all"] = body.published_for_all
    try:
        return db.update_test(test_id, **kwargs)
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.post("/api/tests/{test_id}/replace")
async def api_replace_test(
    test_id: str,
    file: UploadFile = File(...),
    title: Optional[str] = Form(None),
    timer_minutes: Optional[str] = Form(None),
    master_password: Optional[str] = Form(None),
    published_for_all: Optional[str] = Form(None),
    user: dict = Depends(require_user),
):
    """Заменить HTML/банк существующего теста (ссылка и курсанты сохраняются)."""
    t = _require_test_full_editor(test_id, user)
    raw = await file.read()
    try:
        html = raw.decode("utf-8")
    except UnicodeDecodeError:
        html = raw.decode("cp1251", errors="replace")
    if "<html" not in html.lower():
        raise HTTPException(400, "Нужен HTML-файл теста")

    name = (title or "").strip() or (t.get("title") or "Тест")
    m = re.search(r"<title>(.*?)</title>", html, re.I | re.S)
    if m and not (title or "").strip():
        name = re.sub(r"\s*[—\-].*$", "", m.group(1)).strip() or name

    mins = t.get("timer_minutes") or 40
    if timer_minutes is not None and str(timer_minutes).strip() != "":
        try:
            mins = int(str(timer_minutes).strip())
        except ValueError:
            raise HTTPException(400, "Таймер: укажите целое число минут")
        if mins < 1 or mins > 600:
            raise HTTPException(400, "Таймер: от 1 до 600 минут")

    kwargs = {"title": name, "timer_minutes": mins}
    if user.get("role") == "admin" and master_password is not None:
        kwargs["master_password"] = (master_password or "").strip()
    if published_for_all is not None:
        want_pub = _parse_bool_form(published_for_all)
        if want_pub and not _can_publish_for_all(user):
            raise HTTPException(
                403,
                "Публиковать тест для всех могут только администратор и старший инструктор",
            )
        if _can_publish_for_all(user):
            kwargs["published_for_all"] = want_pub

    html = inject_bridge(html, test_id)
    path = os.path.join(db.TESTS_DIR, f"{test_id}.html")
    with open(path, "w", encoding="utf-8") as f:
        f.write(html)
    qn = db.questions_count_from_html(html)
    if qn is not None:
        kwargs["questions_count"] = qn
    try:
        updated = db.update_test(test_id, **kwargs)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {
        "id": test_id,
        "title": updated.get("title") or name,
        "link": f"/t/{test_id}",
        "require_roster": True,
        "timer_minutes": updated.get("timer_minutes") or mins,
        "questions_count": updated.get("questions_count") or qn,
        "has_master_password": bool(updated.get("has_master_password") or updated.get("master_password")),
        "published_for_all": bool(updated.get("published_for_all")),
        "replaced": True,
    }


@app.delete("/api/tests/{test_id}")
def api_delete_test(test_id: str, user: dict = Depends(require_user)):
    require_test_access(test_id, user)
    # инструктор/ассистент — только свои; старший/админ — по доступу
    t = db.get_test(test_id)
    if user.get("role") in ("instructor", "assistant") and t.get("owner_user_id") != user["id"]:
        raise HTTPException(403, "Можно удалять только свои тесты")
    db.delete_test(test_id)
    return {"ok": True}


@app.get("/api/tests/{test_id}/access")
def api_access_info(test_id: str):
    st = db.access_status(test_id)
    if not st:
        raise HTTPException(404, "Тест не найден")
    return st


@app.put("/api/admin/tests/{test_id}/timer")
def api_set_timer(test_id: str, body: TimerIn, user: dict = Depends(require_user)):
    require_test_access(test_id, user)
    try:
        mins = db.set_timer_minutes(test_id, body.minutes)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "timer_minutes": mins, "access": db.access_status(test_id)}


def _master_password_ok(test: dict, pwd: str) -> bool:
    """Пароль теста, иначе глобальный MASTER_PASSWORD (env / admin)."""
    entered = (pwd or "").strip()
    if not entered:
        return False
    stored = (test.get("_master_password") or test.get("master_password") or "").strip()
    if stored:
        return entered == stored
    return entered == MASTER_PASSWORD


@app.post("/api/tests/{test_id}/redeem")
def api_redeem(test_id: str, body: RedeemIn):
    """Мастер-пароль или одноразовый пароль пересдачи."""
    t = db.get_test(test_id)
    if not t:
        raise HTTPException(404, "Тест не найден")

    if body.master_password and _master_password_ok(t, body.master_password):
        token = db.create_session(test_id, is_master=True)
        return {"ok": True, "session": token, "is_master": True}
    if body.master_password:
        raise HTTPException(403, "Неверный мастер-пароль")

    code = (body.code or "").strip()
    if not code:
        raise HTTPException(400, "Введите пароль пересдачи или мастер-пароль")

    inv = db.get_invite_by_code(test_id, code)
    if not inv or not inv.get("for_num"):
        raise HTTPException(403, "Неверный пароль пересдачи")
    if inv["status"] != "unused":
        raise HTTPException(403, "Этот пароль пересдачи уже использован")

    token = db.create_session(
        test_id, invite_id=inv["id"], is_master=False, for_num=inv["for_num"]
    )
    return {
        "ok": True,
        "session": token,
        "is_master": False,
        "for_num": inv["for_num"],
        "code": inv["code"],
    }


def _check_student_access(
    test_id: str, num: str, sess, call: str = "", for_submit: bool = False, allow_started: bool = False
):
    """Проверка списка + однократности. Возвращает (invite_id|None, is_master)."""
    st = db.access_status(test_id)
    is_master = bool(sess and sess.get("is_master") and sess.get("test_id") == test_id)
    if is_master:
        return None, True

    invite_id = None
    retake_ok = False
    if sess and sess.get("test_id") == test_id and sess.get("invite_id"):
        inv = None
        with db.connect() as conn:
            row = conn.execute(
                "SELECT * FROM invites WHERE id=?", (sess["invite_id"],)
            ).fetchone()
            inv = dict(row) if row else None
        if inv and inv["status"] == "unused" and inv.get("for_num"):
            if db._norm_num(num) == db._norm_num(inv["for_num"]):
                invite_id = inv["id"]
                retake_ok = True
            else:
                raise HTTPException(
                    403,
                    "Пароль пересдачи выдан для другого номера: " + str(inv["for_num"]),
                )

    if st and st.get("require_roster"):
        if db.roster_count(test_id) == 0:
            raise HTTPException(
                403,
                "Список допущенных ещё не загружен преподавателем.",
            )
        if not db.in_roster(test_id, num):
            raise HTTPException(
                403,
                "Номер не найден в списке допущенных. Доступ запрещён.",
            )
        if not db.roster_callsign_ok(test_id, num, call or ""):
            raise HTTPException(
                403,
                "Позывной не совпадает со списком допущенных.",
            )

    allowed, att = db.can_start(test_id, num)
    if not allowed and not retake_ok:
        if not (allow_started and att and att.get("status") == "started"):
            raise HTTPException(
                403,
                "Этот номер уже использован. Для пересдачи нужен пароль от преподавателя.",
            )
    return invite_id, False


@app.get("/api/tests/{test_id}/can-start")
def api_can_start(
    test_id: str,
    num: str = Query(...),
    call: str = Query(...),
    x_lis_session: Optional[str] = Header(None),
):
    if not db.get_test(test_id):
        raise HTTPException(404, "Тест не найден")
    sess = db.get_session(x_lis_session) if x_lis_session else None
    try:
        _, is_master = _check_student_access(test_id, num, sess, call=call)
    except HTTPException as e:
        # для UI отдаём структуру с reason
        out = {"allowed": False, "reason": e.detail, "need_retake": "пересдач" in str(e.detail).lower() or "использован" in str(e.detail).lower()}
        if e.status_code in (401, 403):
            return out
        raise

    allowed, att = db.can_start(test_id, num)
    # если есть пароль пересдачи в сессии — разрешаем даже после finished (unlock уже мог быть)
    if not allowed and sess and sess.get("invite_id"):
        allowed = True
    out = {"allowed": allowed, "is_master": is_master}
    if att and att.get("status") == "continue" and allowed:
        raw = att.get("answers_json")
        answers = []
        if raw:
            try:
                answers = json.loads(raw)
            except Exception:
                answers = []
        unanswered = [
            int(a.get("i"))
            for a in answers
            if a.get("pick") is None and a.get("i") is not None
        ]
        out["continue_mode"] = True
        out["continue_minutes"] = att.get("continue_minutes")
        out["unanswered_indices"] = unanswered
        out["previous_answers"] = answers
    if not allowed:
        out["need_retake"] = True
    return out


@app.post("/api/tests/{test_id}/begin")
def api_begin(
    test_id: str,
    body: BeginIn,
    x_lis_session: Optional[str] = Header(None),
):
    if not db.get_test(test_id):
        raise HTTPException(404, "Тест не найден")
    att = db.get_attempt(test_id, body.num)
    sid = (body.start_id or "").strip()
    if att and att.get("status") == "started" and sid and (att.get("start_id") or "") == sid:
        return {"ok": True, "already": True}
    sess = db.get_session(x_lis_session) if x_lis_session else None
    invite_id, is_master = _check_student_access(test_id, body.num, sess, call=body.call)
    if is_master:
        return {"ok": True, "is_master": True}
    try:
        db.begin_attempt(test_id, body.num, body.call, sid, invite_id=invite_id)
    except PermissionError as e:
        raise HTTPException(403, str(e))
    return {"ok": True}


def _answer_picks(items):
    out = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        out.append((item.get("i"), item.get("pick")))
    out.sort(key=lambda pair: (pair[0] is None, pair[0] if isinstance(pair[0], int) else 0))
    return out


def _same_finished_submit(att: Optional[dict], body: SubmitIn) -> bool:
    """Повтор той же сдачи после обрыва ответа, без перезаписи результата."""
    if not att or att.get("status") != "finished":
        return False
    sid = (body.submit_id or "").strip()
    stored_sid = (att.get("submit_id") or "").strip()
    if sid and stored_sid:
        return sid == stored_sid
    try:
        stored = json.loads(att.get("answers_json") or "[]")
    except Exception:
        return False
    picks = _answer_picks(stored)
    return bool(picks) and picks == _answer_picks(body.answers)


@app.post("/api/tests/{test_id}/submit")
def api_submit(
    test_id: str,
    body: SubmitIn,
    x_lis_session: Optional[str] = Header(None),
):
    if not db.get_test(test_id):
        raise HTTPException(404, "Тест не найден")
    existing = db.get_attempt(test_id, body.num)
    if _same_finished_submit(existing, body):
        return {"ok": True, "already": True, "attempt_id": existing["id"]}
    sess = db.get_session(x_lis_session) if x_lis_session else None
    invite_id, is_master = _check_student_access(
        test_id, body.num, sess, call=body.call, for_submit=True, allow_started=True
    )

    att = db.get_attempt(test_id, body.num)
    if att and att.get("status") == "started":
        expected = (att.get("start_id") or "").strip()
        got = (body.start_id or "").strip()
        if not expected or got != expected:
            raise HTTPException(
                403,
                "Этот номер уже использован. Для пересдачи нужен пароль от преподавателя.",
            )
    is_continue = bool(body.is_continue) or (
        att is not None and att.get("status") == "continue"
    )

    # если пересдача — сначала unlock если ещё finished
    if invite_id and not is_continue:
        if att and att["status"] == "finished":
            db.unlock_attempt(test_id, body.num, issue_retake=False)

    try:
        aid = db.upsert_finished_attempt(
            test_id,
            body.num,
            body.call,
            body.score,
            body.total,
            body.percent,
            body.grade,
            body.answers,
            invite_id=None if is_master else invite_id,
            timed_out=bool(body.forced),
            duration_sec=body.duration_sec,
            is_continue=is_continue,
            submit_id=body.submit_id,
        )
    except PermissionError as e:
        raise HTTPException(403, str(e))

    if invite_id and not is_master:
        db.consume_invite(invite_id, body.num, body.call)

    return {"ok": True, "attempt_id": aid, "is_master": is_master, "attempt": db.get_attempt(test_id, body.num)}


# ---------- Roster / retake (admin) ----------
def _roster_entries_from_body(body: RosterIn):
    if body.entries:
        return [
            {
                "num": e.num,
                "call": e.call,
                "fio": e.fio or e.full_name or "",
                "group": e.group_num or e.group or "",
                "stream": e.stream_num or e.stream or "",
            }
            for e in body.entries
        ]
    if body.numbers:
        return [
            {"num": n, "call": "", "fio": "", "group": "", "stream": ""}
            for n in body.numbers
        ]
    return db.parse_roster_text(body.text)


def _cell_to_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v).strip()


def _build_roster_template_bytes():
    """Тот же файл, что и шаблон модуля Курсанты (server/excel_templates.py)."""
    from .excel_templates import build_cadets_template_bytes
    return build_cadets_template_bytes()


def _parse_roster_excel(raw: bytes):
    from .excel_templates import parse_cadets_sheet
    return parse_cadets_sheet(raw)




@app.get("/api/admin/roster-template.xlsx")
def api_roster_template(_: dict = Depends(require_user)):
    data = _build_roster_template_bytes()
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="lis-cadets-template.xlsx"'
        },
    )


@app.get("/api/admin/tests/{test_id}/roster")
def api_get_roster(test_id: str, user: dict = Depends(require_user)):
    require_test_access(test_id, user)
    return {
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
        "retakes": db.list_retake_codes(test_id),
    }


@app.put("/api/admin/tests/{test_id}/roster")
def api_put_roster(test_id: str, body: RosterIn, user: dict = Depends(require_user)):
    require_test_access(test_id, user)
    entries = _roster_entries_from_body(body)
    saved = db.set_roster(test_id, entries)
    return {
        "ok": True,
        "count": len(saved),
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
    }


@app.post("/api/admin/tests/{test_id}/roster/entry")
def api_add_roster_entry(
    test_id: str, body: RosterEntryIn, user: dict = Depends(require_user)
):
    require_test_access(test_id, user)
    num = db._norm_num(body.num)
    if not num:
        raise HTTPException(400, "Укажите номер курсанта")
    if db.get_roster_entry(test_id, num):
        dup = {
            "num": num,
            "call": db._norm_call(body.call or ""),
            "fio": db._norm_fio(body.fio or body.full_name or ""),
            "group": db._norm_call(body.group_num or body.group or ""),
            "stream": db._norm_call(body.stream_num or body.stream or ""),
            "reason": "Запись с таким номером уже существует, проверьте правильность ввода данных",
        }
        return {
            "ok": False,
            "duplicate": True,
            "duplicates": [dup],
            "access": db.access_status(test_id),
            "roster": db.list_roster(test_id),
        }
    try:
        db.upsert_roster_entry(
            test_id,
            body.num,
            body.call,
            body.fio or body.full_name or "",
            group_num=body.group_num or body.group or "",
            stream_num=body.stream_num or body.stream or "",
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {
        "ok": True,
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
    }


@app.delete("/api/admin/tests/{test_id}/roster/{num}")
def api_del_roster_entry(
    test_id: str, num: str, user: dict = Depends(require_user)
):
    require_test_access(test_id, user)
    if not db.remove_roster_entry(test_id, num):
        raise HTTPException(404, "Номер не найден в списке")
    return {
        "ok": True,
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
    }


@app.post("/api/admin/tests/{test_id}/roster/bulk-delete")
def api_bulk_del_roster(
    test_id: str, body: RosterBulkDeleteIn, user: dict = Depends(require_user)
):
    require_test_access(test_id, user)
    nums = [str(x).strip() for x in (body.nums or []) if str(x).strip()]
    if not nums:
        raise HTTPException(400, "Не выбраны записи для удаления")
    deleted = db.remove_roster_entries(test_id, nums)
    return {
        "ok": True,
        "deleted": deleted,
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
    }


@app.post("/api/admin/tests/{test_id}/roster/import")
async def api_import_roster(
    test_id: str,
    mode: str = Form("replace"),
    user: dict = Depends(require_user),
    files: Optional[List[UploadFile]] = File(None),
    file: Optional[UploadFile] = File(None),
):
    """Импорт одного или нескольких Excel. Номер курсанта уникален — дубли не добавляются."""
    require_test_access(test_id, user)
    uploads: List[UploadFile] = []
    if files:
        uploads.extend([f for f in files if f is not None])
    if file is not None:
        uploads.append(file)
    if not uploads:
        raise HTTPException(400, "Выберите файл(ы) Excel (.xlsx)")

    reason = "Запись с таким номером уже существует, проверьте правильность ввода данных"
    duplicates = []
    merged = []  # уникальные записи из всех файлов (первый выигрывает)
    seen_nums = set()
    files_ok = 0

    for up in uploads:
        fname = up.filename or ""
        name = fname.lower()
        if not (name.endswith(".xlsx") or name.endswith(".xlsm")):
            raise HTTPException(400, "Нужен файл Excel (.xlsx): " + fname)
        raw = await up.read()
        try:
            parsed = _parse_roster_excel(raw)
            if isinstance(parsed, tuple):
                entries, file_dups = parsed
            else:
                entries, file_dups = parsed, []
        except Exception as e:
            raise HTTPException(400, "Не удалось прочитать «" + fname + "»: " + str(e))

        if not entries and not file_dups:
            continue
        files_ok += 1

        for d in file_dups:
            item = dict(d)
            item["reason"] = reason
            item["source"] = fname
            duplicates.append(item)

        for e in entries:
            if e["num"] in seen_nums:
                item = dict(e)
                item["reason"] = reason
                item["source"] = fname
                duplicates.append(item)
            else:
                seen_nums.add(e["num"])
                merged.append(dict(e, source=fname))

    if not merged and not duplicates:
        raise HTTPException(400, "В файлах нет строк с номерами")

    if mode == "append":
        existing_rows = db.list_roster(test_id)
        existing_nums = {r["student_num"] for r in existing_rows}
        existing = [
            {
                "num": r["student_num"],
                "call": r.get("roster_callsign") or r.get("callsign") or "",
                "fio": r.get("full_name") or "",
                "group": r.get("group_num") or "",
                "stream": r.get("stream_num") or "",
            }
            for r in existing_rows
        ]
        to_add = []
        for e in merged:
            if e["num"] in existing_nums:
                item = {k: e.get(k) for k in ("num", "call", "fio", "group", "stream")}
                item["reason"] = reason
                item["source"] = e.get("source") or ""
                duplicates.append(item)
            else:
                to_add.append({k: e.get(k) for k in ("num", "call", "fio", "group", "stream")})
                existing_nums.add(e["num"])
        saved = db.set_roster(test_id, existing + to_add)
        imported = len(to_add)
    else:
        # заменить список целиком объединёнными уникальными записями из всех файлов
        to_save = [{k: e.get(k) for k in ("num", "call", "fio", "group", "stream")} for e in merged]
        saved = db.set_roster(test_id, to_save)
        imported = len(to_save)

    return {
        "ok": True,
        "count": len(saved),
        "imported": imported,
        "skipped": len(duplicates),
        "files": files_ok,
        "duplicates": duplicates,
        "access": db.access_status(test_id),
        "roster": db.list_roster(test_id),
    }


@app.post("/api/admin/tests/{test_id}/retake")
def api_retake_code(test_id: str, body: RetakeIn, user: dict = Depends(require_user)):
    require_test_access(test_id, user)
    res = db.unlock_attempt(test_id, body.num, issue_retake=True)
    if not res:
        code = db.create_retake_code(test_id, body.num)
        return {"ok": True, "unlocked": False, "invite": code, "retake": code}
    return res


@app.get("/api/admin/attempts")
def api_attempts(test_id: Optional[str] = None, user: dict = Depends(require_user)):
    ids = auth.visible_test_ids(user)
    rows = db.list_attempts(test_id)
    if ids is None:
        return rows
    idset = set(ids)
    if test_id and test_id not in idset:
        return []
    return [r for r in rows if r.get("test_id") in idset]


@app.get("/api/admin/attempts/{attempt_id}")
def api_attempt_detail(attempt_id: str, user: dict = Depends(require_user)):
    att = db.get_attempt_by_id(attempt_id)
    if not att:
        raise HTTPException(404, "Попытка не найдена")
    require_test_access(att["test_id"], user)
    return att


@app.post("/api/admin/attempts/{attempt_id}/continue")
def api_allow_continue(
    attempt_id: str, body: ContinueIn, user: dict = Depends(require_user)
):
    att = db.get_attempt_by_id(attempt_id)
    if not att:
        raise HTTPException(404, "Попытка не найдена")
    require_test_access(att["test_id"], user)
    try:
        return {"ok": True, "attempt": db.allow_continue_attempt(attempt_id, body.minutes)}
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.put("/api/admin/attempts/{attempt_id}/grade")
def api_set_attempt_grade(
    attempt_id: str, body: GradeIn, user: dict = Depends(require_user)
):
    if user.get("role") not in ("admin", "senior"):
        raise HTTPException(
            403, "Менять оценку могут только администратор и старший инструктор"
        )
    att = db.get_attempt_by_id(attempt_id)
    if not att:
        raise HTTPException(404, "Попытка не найдена")
    require_test_access(att["test_id"], user)
    try:
        return {"ok": True, "attempt": db.set_attempt_grade(attempt_id, body.grade)}
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.post("/api/admin/unlock")
def api_unlock(body: UnlockIn, user: dict = Depends(require_user)):
    require_test_access(body.test_id, user)
    res = db.unlock_attempt(
        body.test_id, body.num, body.call, issue_retake=body.issue_invite
    )
    if not res:
        raise HTTPException(404, "Завершённая попытка не найдена")
    return res


def _cadets_for_user(user: dict):
    """Все видят полный список; флаги can_edit / can_delete — по владельцу или admin."""
    rows = db.list_cadets()
    out = []
    for c in rows:
        d = dict(c)
        d["can_edit"] = auth.can_modify_cadet(user, c)
        d["can_delete"] = auth.can_modify_cadet(user, c)
        out.append(d)
    return out


def _build_cadets_template_bytes():
    """Шаблон курсантов: server/excel_templates.py, список CADET_COLUMNS."""
    from .excel_templates import build_cadets_template_bytes
    return build_cadets_template_bytes()


def _parse_cadets_excel(raw: bytes):
    from .excel_templates import parse_cadets_sheet
    entries, _dups = parse_cadets_sheet(raw)
    return [
        {
            "student_num": e["num"],
            "callsign": e.get("call") or "",
            "group_num": e.get("group") or "",
            "stream_num": e.get("stream") or "",
            "full_name": e.get("fio") or "",
        }
        for e in entries
    ]


# ---------- Cadets ----------
@app.get("/api/admin/cadets")
def api_list_cadets(user: dict = Depends(require_user)):
    return {"cadets": _cadets_for_user(user)}


@app.get("/api/admin/cadets-template.xlsx")
def api_cadets_template(_: dict = Depends(require_user)):
    data = _build_cadets_template_bytes()
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": 'attachment; filename="lis-cadets-template.xlsx"'
        },
    )


@app.post("/api/admin/cadets/import")
async def api_import_cadets(
    file: UploadFile = File(...),
    user: dict = Depends(require_user),
):
    name = (file.filename or "").lower()
    if not (name.endswith(".xlsx") or name.endswith(".xlsm")):
        raise HTTPException(400, "Нужен файл Excel (.xlsx)")
    raw = await file.read()
    try:
        entries = _parse_cadets_excel(raw)
    except Exception as e:
        raise HTTPException(400, "Не удалось прочитать Excel: " + str(e))
    if not entries:
        raise HTTPException(400, "В файле нет строк с номером курсанта")

    report = []
    added = 0
    seen_in_file = set()
    for e in entries:
        num = e["student_num"]
        item = {
            "student_num": num,
            "callsign": e.get("callsign") or "",
            "group_num": e.get("group_num") or "",
            "stream_num": e.get("stream_num") or "",
            "status": "skipped",
            "reason": "",
        }
        if num in seen_in_file:
            item["reason"] = "Номер уже существует"
            report.append(item)
            continue
        seen_in_file.add(num)
        if db.cadet_num_exists(num):
            item["reason"] = "Номер уже существует"
            report.append(item)
            continue
        try:
            db.add_cadet(
                e.get("callsign") or "",
                num,
                e.get("full_name") or "",
                group_num=e.get("group_num") or "",
                stream_num=e.get("stream_num") or "",
                owner_user_id=user["id"],
                department_id=user.get("department_id"),
            )
            item["status"] = "added"
            item["reason"] = ""
            added += 1
        except ValueError as err:
            item["reason"] = str(err)
            if "уже существует" in str(err):
                item["reason"] = "Номер уже существует"
        report.append(item)

    report.sort(key=lambda r: (0 if r["status"] == "skipped" else 1, r["student_num"]))
    return {
        "ok": True,
        "added": added,
        "total": len(entries),
        "rows": report,
        "cadets": _cadets_for_user(user),
    }


@app.post("/api/admin/cadets")
def api_add_cadet(body: CadetIn, user: dict = Depends(require_user)):
    try:
        row = db.add_cadet(
            body.callsign,
            body.student_num,
            body.full_name,
            group_num=body.group_num,
            stream_num=body.stream_num,
            owner_user_id=user["id"],
            department_id=user.get("department_id"),
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "cadet": row, "cadets": _cadets_for_user(user)}


@app.put("/api/admin/cadets/{cadet_id}")
def api_update_cadet(cadet_id: str, body: CadetIn, user: dict = Depends(require_user)):
    cur = db.get_cadet(cadet_id)
    if not cur:
        raise HTTPException(404, "Курсант не найден")
    if not auth.can_modify_cadet(user, cur):
        raise HTTPException(403, "Можно менять только своих курсантов")
    try:
        row = db.update_cadet(
            cadet_id,
            body.callsign,
            body.student_num,
            body.full_name,
            group_num=body.group_num,
            stream_num=body.stream_num,
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    if not row:
        raise HTTPException(404, "Курсант не найден")
    return {"ok": True, "cadet": row, "cadets": _cadets_for_user(user)}


@app.delete("/api/admin/cadets/{cadet_id}")
def api_delete_cadet(cadet_id: str, user: dict = Depends(require_user)):
    cur = db.get_cadet(cadet_id)
    if not cur:
        raise HTTPException(404, "Курсант не найден")
    if not auth.can_modify_cadet(user, cur):
        raise HTTPException(403, "Можно удалять только своих курсантов")
    if not db.remove_cadet(cadet_id):
        raise HTTPException(404, "Курсант не найден")
    return {"ok": True, "cadets": _cadets_for_user(user)}


@app.post("/api/admin/cadets/bulk-delete")
def api_bulk_delete_cadets(body: CadetBulkDeleteIn, user: dict = Depends(require_user)):
    ids = [str(x).strip() for x in (body.ids or []) if str(x).strip()]
    if not ids:
        raise HTTPException(400, "Не выбраны записи для удаления")
    allowed = []
    denied = 0
    for cid in ids:
        cur = db.get_cadet(cid)
        if not cur:
            continue
        if auth.can_modify_cadet(user, cur):
            allowed.append(cid)
        else:
            denied += 1
    if not allowed and denied:
        raise HTTPException(403, "Можно удалять только своих курсантов")
    if not allowed:
        raise HTTPException(400, "Не выбраны записи для удаления")
    deleted = db.remove_cadets(allowed)
    return {
        "ok": True,
        "deleted": deleted,
        "denied": denied,
        "cadets": _cadets_for_user(user),
    }


class FlightSheetRowsIn(BaseModel):
    rows: List[Any] = Field(default_factory=list)
    days: List[Any] = Field(default_factory=list)


def _flight_sheet_for_user(user_id: str):
    """Актуальная ведомость. Старый произвольный шаблон заменяется этой формой."""
    cadets = db.list_cadets()
    current = db.get_flight_sheet(user_id)
    if current and not flight_sheet.is_current_schema(current):
        built = flight_sheet.migrate_sheet(current, cadets)
        current = db.save_flight_sheet(
            user_id, built.get("source_name") or "", built["columns"], built["rows"]
        )
    if current:
        current = flight_sheet.refresh_sheet(current, cadets)
    return current, cadets


def _flight_sheet_edits(current, body):
    current = dict(current)
    if body.rows:
        current["rows"] = flight_sheet.sanitize_rows(current.get("columns") or [], body.rows)
    if body.days:
        current["columns"] = flight_sheet.apply_day_headers(current.get("columns") or [], body.days)
    return current


def _require_flight_sheet_user(user: dict) -> dict:
    if not flight_sheet.user_can_use(user):
        raise HTTPException(403, "Блок доступен учётной записи СИ_Тестер")
    return user


@app.get("/api/admin/flight-sheet")
def api_get_flight_sheet(user: dict = Depends(require_user)):
    _require_flight_sheet_user(user)
    sheet, _cadets = _flight_sheet_for_user(user["id"])
    return {"sheet": sheet}


@app.get("/api/admin/flight-sheet/template.xlsx")
def api_flight_sheet_template(user: dict = Depends(require_user)):
    _require_flight_sheet_user(user)
    current, cadets = _flight_sheet_for_user(user["id"])
    sheet = current or flight_sheet.build_sheet(cadets, days=1)
    data = flight_sheet.build_template_bytes(sheet)
    return Response(
        data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="polet-vedomost.xlsx"'},
    )


@app.post("/api/admin/flight-sheet/day")
def api_add_flight_sheet_day(body: FlightSheetRowsIn, user: dict = Depends(require_user)):
    _require_flight_sheet_user(user)
    current, cadets = _flight_sheet_for_user(user["id"])
    if not current:
        built = flight_sheet.build_sheet(cadets, days=1)
        message = "Таблица собрана: один день. Кнопка «+» добавляет следующий."
    else:
        current = _flight_sheet_edits(current, body)
        try:
            built = flight_sheet.add_day(current, cadets)
        except ValueError as e:
            raise HTTPException(400, str(e))
        last = flight_sheet.headers_from_columns(built["columns"])[-1]
        message = f"Добавлен {last.get('title') or 'день'} · {last.get('date') or ''}"
    saved = db.save_flight_sheet(
        user["id"],
        built.get("source_name") or "",
        built["columns"],
        built["rows"],
    )
    return {"ok": True, "message": message, "sheet": saved}


@app.post("/api/admin/flight-sheet/import")
async def api_import_flight_sheet(
    file: UploadFile = File(...),
    user: dict = Depends(require_user),
):
    _require_flight_sheet_user(user)
    raw = await file.read()
    if not raw:
        raise HTTPException(400, "Файл пустой")
    if len(raw) > 8 * 1024 * 1024:
        raise HTTPException(400, "Файл больше 8 МБ")
    cadets = db.list_cadets()
    try:
        built = flight_sheet.parse_template(raw, file.filename or "", cadets)
    except ValueError as e:
        raise HTTPException(400, str(e))
    except Exception as e:
        raise HTTPException(400, "Не удалось прочитать файл: " + str(e))
    saved = db.save_flight_sheet(
        user["id"],
        built.get("source_name") or "",
        built["columns"],
        built["rows"],
    )
    return {
        "ok": True,
        "message": built.get("message") or "",
        "sheet": saved,
    }


@app.put("/api/admin/flight-sheet")
def api_save_flight_sheet(body: FlightSheetRowsIn, user: dict = Depends(require_user)):
    _require_flight_sheet_user(user)
    current, cadets = _flight_sheet_for_user(user["id"])
    if not current:
        raise HTTPException(400, "Сначала сформируйте таблицу")
    try:
        current = _flight_sheet_edits(current, body)
        built = flight_sheet.refresh_sheet(current, cadets)
        saved = db.save_flight_sheet(
            user["id"],
            built.get("source_name") or "",
            built["columns"],
            built["rows"],
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {"ok": True, "sheet": saved}


@app.delete("/api/admin/flight-sheet")
def api_delete_flight_sheet(user: dict = Depends(require_user)):
    _require_flight_sheet_user(user)
    db.delete_flight_sheet(user["id"])
    return {"ok": True}


def _require_library_user(user: dict) -> dict:
    if not library.user_can_use(user):
        raise HTTPException(403, "Раздел доступен учётной записи СИ_Тестер")
    return user


@app.get("/api/admin/library")
def api_list_library(category: str = Query("materials"), user: dict = Depends(require_user)):
    _require_library_user(user)
    if category not in library.CATEGORIES:
        raise HTTPException(400, "Неизвестный раздел библиотеки")
    docs = []
    for row in db.list_library_docs(category):
        docs.append(
            {
                "id": row["id"],
                "category": row["category"],
                "title": row["title"],
                "summary": row["summary"],
                "filename": row["original_name"],
            }
        )
    return {"docs": docs}


@app.post("/api/admin/library")
async def api_add_library_doc(
    file: UploadFile = File(...),
    category: str = Form(...),
    title: str = Form(""),
    summary: str = Form(""),
    user: dict = Depends(require_user),
):
    _require_library_user(user)
    raw = await file.read()
    try:
        doc = library.save_upload(category, title, summary, file.filename or "", raw)
    except ValueError as e:
        raise HTTPException(400, str(e))
    return {
        "ok": True,
        "doc": {
            "id": doc["id"],
            "category": doc["category"],
            "title": doc["title"],
            "summary": doc["summary"],
            "filename": doc["original_name"],
        },
    }


@app.get("/api/admin/library/{doc_id}/file")
def api_library_file(doc_id: str, user: dict = Depends(require_user)):
    _require_library_user(user)
    doc = db.get_library_doc(doc_id)
    if not doc:
        raise HTTPException(404, "Документ не найден")
    try:
        path = library.file_path(doc)
    except ValueError as e:
        raise HTTPException(404, str(e))
    import mimetypes
    from urllib.parse import quote

    media = mimetypes.guess_type(doc.get("original_name") or "")[0] or "application/octet-stream"
    filename = quote(doc.get("original_name") or "file")
    return FileResponse(
        path,
        media_type=media,
        headers={"Content-Disposition": f"inline; filename*=UTF-8''{filename}"},
    )


@app.delete("/api/admin/library/{doc_id}")
def api_delete_library_doc(doc_id: str, user: dict = Depends(require_user)):
    _require_library_user(user)
    if not library.remove_doc(doc_id):
        raise HTTPException(404, "Документ не найден")
    return {"ok": True}


@app.post("/api/admin/cadets/clear-mine")
def api_clear_my_cadets(user: dict = Depends(require_user)):
    """Очистить курсантов, добавленных текущей учётной записью."""
    uid = user.get("id")
    if not uid:
        raise HTTPException(400, "Не удалось определить учётную запись")
    deleted = db.clear_cadets_for_owner(uid)
    return {"ok": True, "deleted": deleted, "cadets": _cadets_for_user(user)}


# ---------- Pages ----------
_NAV_SNIPPET = (
    '<link rel="stylesheet" href="/static/lis-nav.css">'
    '<script src="/static/lis-nav.js" defer></script>'
)


def _with_nav(html: str) -> str:
    """Подключает кнопки «На главную» / «Назад» на странице."""
    if "lis-nav.js" in html:
        return html
    low = html.lower()
    idx = low.find("</head>")
    if idx >= 0:
        return html[:idx] + _NAV_SNIPPET + html[idx:]
    idx = low.find("<body")
    if idx >= 0:
        end = html.find(">", idx)
        if end >= 0:
            return html[: end + 1] + _NAV_SNIPPET + html[end + 1 :]
    return _NAV_SNIPPET + html


def _read_html(path: str) -> str:
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


def _html_page(path: str, *, nav: bool = False) -> HTMLResponse:
    """HTML с автоподстановкой ?v= по хэшу файлов /static/…."""
    html = _read_html(path)
    if nav:
        html = _with_nav(html)
    html = static_version.bust_static_urls(html, STATIC_DIR)
    return HTMLResponse(html)


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    denied = _guard_html(request)
    if denied:
        return denied
    path = os.path.join(STATIC_DIR, "hub.html")
    return _html_page(path, nav=True)


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request):
    # без «На главную» / «Назад»; уже вошёл — на главную
    if _user_from_request(request):
        return RedirectResponse(url="/", status_code=303)
    path = os.path.join(STATIC_DIR, "login.html")
    return _html_page(path)


@app.get("/admin", response_class=HTMLResponse)
def admin_page(request: Request):
    denied = _guard_html(request)
    if denied:
        return denied
    path = os.path.join(STATIC_DIR, "admin.html")
    return _html_page(path, nav=True)


@app.get("/t/{test_id}")
def take_test(test_id: str):
    if not db.get_test(test_id):
        raise HTTPException(404, "Тест не найден")
    path = os.path.join(db.TESTS_DIR, f"{test_id}.html")
    if not os.path.isfile(path):
        raise HTTPException(404, "Файл теста отсутствует")
    # страница курсанта: без «На главную» / «Назад»
    html = _read_html(path)
    return HTMLResponse(html, media_type="text/html; charset=utf-8")


@app.get("/static/generator/generator-app.js")
def generator_app_js():
    """JS генератора: внутри тоже подставляем ?v= для xlsx/tpl."""
    path = os.path.join(STATIC_DIR, "generator", "generator-app.js")
    if not os.path.isfile(path):
        raise HTTPException(404, "Файл не найден")
    raw = _read_html(path)
    body = static_version.bust_static_urls(raw, STATIC_DIR, strings=True)
    return Response(body, media_type="application/javascript; charset=utf-8")


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")

TOOL_FILES = {
    "generator": "Генератор_тестов.html",
    "converter": "Конвертер_тестов.html",
    "norms": "Нормативы.html",
    "flights": "Полётная_ведомость.html",
    "polygon": "Полигон.html",
}

STATIC_TOOL_FILES = {
    "workload": "workload.html",
}


class WorkloadCellIn(BaseModel):
    year: int = Field(..., ge=2000, le=2100)
    month: int = Field(..., ge=1, le=12)
    staff_id: str = Field(..., min_length=1)
    day: int = Field(..., ge=1, le=31)
    code: str = ""
    hours: str = ""


def _require_feature(module_id: str, user: dict):
    try:
        features.require_module(module_id, user)
    except ValueError as e:
        raise HTTPException(403, str(e))


@app.get("/api/workload")
def api_get_workload(
    year: int = Query(..., ge=2000, le=2100),
    month: int = Query(..., ge=1, le=12),
    user: dict = Depends(require_user),
):
    _require_feature("workload", user)
    return {"year": year, "month": month, "cells": db.get_workload_month(year, month)}


@app.put("/api/workload/cell")
def api_put_workload_cell(body: WorkloadCellIn, user: dict = Depends(require_user)):
    _require_feature("workload", user)
    staff = auth.get_staff(body.staff_id)
    if not staff:
        raise HTTPException(404, "Сотрудник не найден")
    if user.get("role") not in ("admin", "senior"):
        if staff.get("department_id") != user.get("department_id"):
            raise HTTPException(403, "Нет доступа к этому сотруднику")
    try:
        return db.upsert_workload_cell(
            body.year,
            body.month,
            body.staff_id,
            body.day,
            code=body.code,
            hours=body.hours,
        )
    except ValueError as e:
        raise HTTPException(400, str(e))


class WorkloadPolygonIn(BaseModel):
    staff_id: str = Field(..., min_length=1)
    name: str = ""
    date_from: str = ""
    date_to: str = ""


def _require_staff_access(staff_id: str, user: dict) -> dict:
    staff = auth.get_staff(staff_id)
    if not staff:
        raise HTTPException(404, "Сотрудник не найден")
    if user.get("role") not in ("admin", "senior"):
        if staff.get("department_id") != user.get("department_id"):
            raise HTTPException(403, "Нет доступа к этому сотруднику")
    return staff


@app.get("/api/workload/polygon/{staff_id}")
def api_get_workload_polygon(staff_id: str, user: dict = Depends(require_user)):
    _require_feature("workload", user)
    _require_staff_access(staff_id, user)
    return db.get_workload_polygon(staff_id)


@app.put("/api/workload/polygon")
def api_put_workload_polygon(body: WorkloadPolygonIn, user: dict = Depends(require_user)):
    _require_feature("workload", user)
    _require_staff_access(body.staff_id, user)
    try:
        return db.upsert_workload_polygon(
            body.staff_id,
            name=body.name,
            date_from=body.date_from,
            date_to=body.date_to,
        )
    except ValueError as e:
        raise HTTPException(400, str(e))


@app.get("/tools/{name}", response_class=HTMLResponse)
def tool_page(name: str, request: Request):
    denied = _guard_html(request)
    if denied:
        return denied
    feature_id = features.TOOL_FEATURE.get(name)
    if feature_id:
        user = _user_from_request(request)
        if not features.can_access(feature_id, user):
            return RedirectResponse(url="/", status_code=303)
    static_name = STATIC_TOOL_FILES.get(name)
    if static_name:
        path = os.path.join(STATIC_DIR, static_name)
        if not os.path.isfile(path):
            raise HTTPException(404, "Файл не найден")
        return _html_page(path, nav=True)
    fname = TOOL_FILES.get(name)
    if not fname:
        raise HTTPException(404)
    path = os.path.join(ROOT, fname)
    if not os.path.isfile(path):
        raise HTTPException(404, "Файл не найден")
    return _html_page(path, nav=True)
